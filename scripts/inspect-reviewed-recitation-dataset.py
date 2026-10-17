"""Inspect licensed expert-labelled public data without guessing its label schema."""
import collections
import csv
import hashlib
import io
import json
from pathlib import Path
import urllib.request
import urllib.error
import zipfile
import openpyxl

URL = "https://data.mendeley.com/public-files/datasets/sxtmmr6mvk/files/ec1b3a42-700d-4424-b6cd-6844828f4637/file_downloaded"
SHA = "de8042ac8cc137e83fdcdbed375a24bd51f873c94f7f2cca1e39c60b188c0068"
path = Path("reviewed-public-recordings.zip")
digest = hashlib.sha256()
try:
    with urllib.request.urlopen(URL, timeout=180) as response, path.open("wb") as stream:
        while chunk := response.read(1024 * 1024):
            digest.update(chunk)
            stream.write(chunk)
except urllib.error.HTTPError as error:
    if error.code != 403:
        raise
    revision = "973b5ef1d1532bcbe7e9a5a5b0dce125cc4d0dac"
    mirror = f"https://huggingface.co/datasets/MuazAhmad7/Surah_Ikhlas-Labeled_Dataset/resolve/{revision}/metadata.csv"
    with urllib.request.urlopen(mirror, timeout=60) as response:
        raw = response.read(2_000_001)
    if len(raw) > 2_000_000:
        raise ValueError("Unexpected label metadata length")
    rows = list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))
    report = dict(originalArchiveAccessHTTP=403, originalArchiveSHA256Verified=False,
        originalDOI="10.17632/sxtmmr6mvk.1", mirrorRepository="MuazAhmad7/Surah_Ikhlas-Labeled_Dataset",
        mirrorRevision=revision, mirrorLicense="CC-BY-4.0", mirrorMetadataSHA256=hashlib.sha256(raw).hexdigest(),
        rows=len(rows), columns=list(rows[0]), labels=dict(collections.Counter(r["label_name"] for r in rows)),
        lineage="Published mirror labels inspected; equivalence to the original expert-reviewed archive is NOT established")
    Path("reviewed-dataset-schema.json").write_text(json.dumps(report, ensure_ascii=False, indent=2)+"\n")
    print("Original archive access blocked (HTTP403); accessible published label schema inspected, expert lineage remains unverified.")
    raise SystemExit(0)
if digest.hexdigest() != SHA:
    raise ValueError("Published dataset checksum mismatch")
report = dict(doi="10.17632/sxtmmr6mvk.1", license="CC-BY-4.0", sha256=SHA,
              citation="Maghraby, Ashwag; Alsaedi, Manar; Mahjoob, Azzah; Al-Otaibi, Amirah; Alsaedi, Waad; Alsaedi, Ather (2024), Surah Al-Ikhlas of the Holy Qur'an Error Detection Dataset, Mendeley Data, V1",
              scope="Schema inspection only; no classifier accuracy claimed", entries=[], spreadsheets=[])
with zipfile.ZipFile(path) as archive:
    report["entries"] = [dict(name=i.filename, bytes=i.file_size) for i in archive.infolist() if not i.is_dir()]
    report["extensions"] = dict(collections.Counter(Path(i.filename).suffix.lower() for i in archive.infolist() if not i.is_dir()))
    for item in archive.infolist():
        if item.filename.lower().endswith(".xlsx") and item.file_size <= 8_000_000:
            workbook = openpyxl.load_workbook(io.BytesIO(archive.read(item)), read_only=True, data_only=True)
            document = dict(name=item.filename, sheets=[])
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                # Observe the schema and finite categorical value counts; don't
                # replicate free text, personal contact fields or the audio.
                counts = [collections.Counter() for _ in header]
                for row in rows:
                    for index, value in enumerate(row[:len(counts)]):
                        if value is not None:
                            counts[index][str(value)] += 1
                document["sheets"].append(dict(name=sheet.title, rows=sheet.max_row,
                    columns=[dict(header=str(h), unique=len(c), categories=dict(c) if len(c) <= 20 and all(len(v) < 100 for v in c) else None)
                             for h, c in zip(header, counts)]))
            workbook.close()
            report["spreadsheets"].append(document)
Path("reviewed-dataset-schema.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
print("Verified public dataset archive and inspected its actual label schema.")
